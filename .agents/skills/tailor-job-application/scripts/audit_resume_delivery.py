#!/usr/bin/env python3
"""Audit resume content, evidence boundaries, and confirmed DOCX formatting."""

from __future__ import annotations

import argparse
import os
import re
import sys
import zipfile
import json
import hashlib
from difflib import SequenceMatcher
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

W_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
FONT_NAME = "Microsoft YaHei"
NAME_MARKERS = ()
CASE_PATTERN = r"(?!)"
FORBIDDEN_PATTERNS = (
    ("建议改为", r"建\s*议改为"),
    ("缺少证据", r"缺少\s*证据"),
    ("证据等级", r"证据\s*等级"),
    ("匹配度", r"匹配\s*度"),
    ("置信度", r"置信\s*度"),
    ("Apply–P", r"Apply\s*[–-]\s*P"),
    ("Pending", r"\bPending\b"),
    ("Skip", r"\bSkip\b"),
    ("风险提示", r"风险\s*提示"),
    ("内部判断", r"内部\s*判断"),
    ("AI判断", r"AI\s*判断"),
    ("草稿备注", r"草稿\s*备注"),
    ("占位或待核验", r"TODO|TBD|待核验|待确认|待补充|占位符|\{\{.+?\}\}"),
)


def effective_chars(text: str) -> int:
    return sum(1 for char in text if not char.isspace())


def has_page_break(paragraph) -> bool:
    return any(br.get(qn("w:type")) == "page" for br in paragraph._p.iter(qn("w:br")))


def has_section_break(paragraph) -> bool:
    p_pr = paragraph._p.pPr
    return p_pr is not None and p_pr.sectPr is not None


def find_resume_divider(doc: Document) -> int:
    paragraphs = doc.paragraphs
    for index, paragraph in enumerate(paragraphs):
        if not (has_page_break(paragraph) or has_section_break(paragraph)):
            continue
        siblings = list(paragraph._p.itersiblings())[:4]
        following = " ".join("".join(n.text or "" for n in item.iter(qn("w:t"))) for item in siblings)
        if any(marker.lower() in following.lower() for marker in NAME_MARKERS):
            return index
    raise ValueError("未找到审阅区与简历正文之间的可验证分隔")


def document_text(path: Path, after_page_break: bool = False) -> str:
    if path.suffix.lower() != ".docx":
        raise ValueError(f"只支持 .docx：{path}")
    doc = Document(path)
    start = find_resume_divider(doc) + 1 if after_page_break else 0
    # Include tables/text boxes in document order, including the name/photo table.
    boundary = doc.paragraphs[start - 1]._p if start else None
    active = start == 0
    lines = []
    for p in doc.element.body.iter(qn("w:p")):
        if p is boundary:
            active = True
            continue
        if active:
            value = "".join(n.text or "" for n in p.iter(qn("w:t")))
            if value.strip(): lines.append(value.strip())
    return "\n".join(lines)


def normalized(text: str) -> str:
    return re.sub(r"\s+", "", text)


def source_text(path: Path) -> str:
    if path.suffix.lower() == ".docx":
        return document_text(path)
    if path.suffix.lower() == ".pdf":
        from pypdf import PdfReader
        return "\n".join(p.extract_text() or "" for p in PdfReader(path).pages)
    return path.read_text(encoding="utf-8-sig")


def context_errors(text: str, context: dict) -> list[str]:
    """Check actual output against the existing per-job evidence/edit record.

    This does not infer semantic truth. Source selection, completeness, and
    narrative fidelity still need a paragraph-by-paragraph review.
    All paths are read-only. No scoring or new mandatory worksheet is introduced.
    """
    errors = []
    norm = normalized(text)
    sources = {}
    for key, item in context.get("sources", {}).items():
        p = Path(item["path"])
        if not p.is_file():
            errors.append(f"来源不可读:{key}"); continue
        actual = hashlib.sha256(p.read_bytes()).hexdigest()
        if actual != item.get("sha256"):
            errors.append(f"来源版本已变:{key}"); continue
        sources[key] = normalized(source_text(p))

    def cited(ref):
        quote = normalized(ref.get("quote", ""))
        if not quote or quote not in sources.get(ref.get("source"), ""):
            errors.append(f"来源原段未对应:{ref.get('source')}")
        return quote

    for lock in context.get("locked_passages", []):
        expected = cited(lock)
        if expected and expected not in norm:
            errors.append("锁定段落被改动或丢失:" + expected[:35])
    for token in context.get("foreign_targets", []):
        if token and normalized(token) in norm:
            errors.append("其他岗位残留:" + token)
    for token in context.get("required_target", []):
        if normalized(token) not in normalized("\n".join(text.splitlines()[:10])):
            errors.append("目标岗位抬头不符:" + token)

    blocks = context.get("blocks", [])
    positions = []
    for block in blocks:
        anchor = normalized(block["anchor"])
        pos = norm.find(anchor)
        if pos < 0: errors.append("核心项目/段落缺失:" + block["anchor"])
        positions.append(pos)
    for block, start in zip(blocks, positions):
        if start < 0: continue
        end = min([p for p in positions if p > start] + [len(norm)])
        # Optional next heading prevents scanning unrelated projects/employers.
        if block.get("end_before"):
            limit = norm.find(normalized(block["end_before"]), start + len(normalized(block["anchor"])))
            if limit >= 0: end = min(end, limit)
            else: errors.append("项目边界缺失:" + block["end_before"])
        actual = norm[start:end]
        refs = "".join(cited(r) for r in block.get("sources", []))
        employer = block.get("employer")
        if employer:
            # Only dated employment headings define employment; mentions in prose do not.
            headings = [(m.start(),m.group()) for m in re.finditer(r"(?m)^\s*20\d\d[.年/-].*", text)]
            raw_start = text.find(block["anchor"])
            before = [h for p,h in headings if p < raw_start]
            if not before or employer not in before[-1]:
                errors.append("项目公司归属不符:" + block["anchor"])
        for group in block.get("required_evidence", []):
            if not any(normalized(choice) in actual for choice in group):
                errors.append("证据关系被删薄:" + block["anchor"] + ":" + "/".join(group))
        for pattern in block.get("forbidden_patterns", []):
            if re.search(pattern, actual, re.I):
                errors.append("职责或事实边界越界:" + block["anchor"] + ":" + pattern)
        if block.get("check_numbers"):
            numeric_pattern = r"\d+(?:\.\d+)?(?:[%％]|万|亿)?"
            source_numbers = set(re.findall(numeric_pattern, refs))
            for number in set(re.findall(numeric_pattern, actual)):
                if number not in source_numbers:
                    errors.append("项目数字未见于对应来源:" + block["anchor"] + ":" + number)

    # Local edits: compare the whole before/after sequence, not just the edited paragraph.
    scope = context.get("local_edit")
    if scope:
        before = scope["before_paragraphs"]
        after = [p for p in text.splitlines() if p.strip()]
        allowed = set(scope["allowed_before_indices"])
        for op,a,b,c,d in SequenceMatcher(None,before,after,autojunk=False).get_opcodes():
            if op == "equal": continue
            impacted = set(range(a,b)) if a != b else {a-1,a}
            if not impacted.intersection(allowed) or (a != b and not impacted.issubset(allowed)):
                errors.append(f"局部修改越界:{op}:原段{a}:{b}")
    return errors


def revision_errors(path: Path) -> list[str]:
    from xml.etree import ElementTree as ET
    with zipfile.ZipFile(path) as z:
        root = ET.fromstring(z.read("word/document.xml"))
        return ["正文存在未处理修订/批注:" + tag for tag in ("ins","del","commentRangeStart")
                if root.find('.//' + W_NS + tag) is not None]


def summary_errors(text: str) -> list[str]:
    """Reject concrete case narration in this candidate's opening summary."""
    lines = text.splitlines()
    start = next((i for i, line in enumerate(lines)
                  if re.sub(r"[\s|｜]", "", line) in ("个人简介", "职业摘要", "核心优势")), None)
    if start is None:
        return ["未定位个人简介，无法完成首页写法检查；须检查真实正文结构"]
    end = next((i for i in range(start + 1, len(lines))
                if re.sub(r"[\s|｜]", "", lines[i]) in ("工作经历", "工作经历及项目", "教育背景", "教育经历")), len(lines))
    summary = "\n".join(lines[start + 1:end])
    patterns = (
        ("具体案例名称", CASE_PATTERN),
        ("案例业绩数字", r"(?:GMV|CTR|CPM|ROI)\s*[达为约：:]?\s*\d|\d+(?:\.\d+)?\s*(?:万|亿|%|％)|\d+\s*位.{0,4}达人|\d+\+?\s*(?:家)?门店"),
        ("举例式案例展开", r"例如|以.{1,20}为例|在.{1,20}项目中"),
    )
    return ["个人简介含" + label + "，应改为JD核心能力陈述，案例放经历正文"
            for label, pattern in patterns if re.search(pattern, summary, flags=re.IGNORECASE)]


def font_installed() -> bool:
    if os.name != "nt":
        return True
    fonts = Path(os.environ.get("WINDIR", r"C:\Windows")) / "Fonts"
    return any((fonts / name).exists() for name in ("msyh.ttc", "msyhbd.ttc", "msyhl.ttc"))


def run_font_errors(docx_path: Path) -> list[str]:
    errors: list[str] = []
    with zipfile.ZipFile(docx_path) as archive:
        parts = [name for name in archive.namelist() if name == "word/document.xml" or re.fullmatch(r"word/(header|footer)\d+\.xml", name)]
        from xml.etree import ElementTree as ET

        for part in parts:
            root = ET.fromstring(archive.read(part))
            for run_index, run in enumerate(root.iter(W_NS + "r"), start=1):
                visible = "".join(node.text or "" for node in run.iter(W_NS + "t")).strip()
                if not visible:
                    continue
                r_pr = run.find(W_NS + "rPr")
                r_fonts = r_pr.find(W_NS + "rFonts") if r_pr is not None else None
                values = {r_fonts.get(W_NS + attr) if r_fonts is not None else None for attr in ("ascii", "hAnsi", "eastAsia", "cs")}
                if values != {FONT_NAME}:
                    errors.append(f"{part}:run{run_index}:未显式统一为微软雅黑:{visible[:30]}")
                    if len(errors) >= 20:
                        return errors
    return errors


def twips(length) -> int:
    return round(length.inches * 1440)


def format_errors(path: Path, after_page_break: bool) -> list[str]:
    doc = Document(path)
    errors: list[str] = []
    errors.extend(run_font_errors(path))
    if after_page_break:
        try:
            divider = find_resume_divider(doc)
        except ValueError as exc:
            return errors + [str(exc)]
        if len(doc.sections) < 2:
            errors.append("审阅区与简历正文未使用独立分节")
        body_paragraphs = doc.paragraphs[divider + 1 :]
    else:
        body_paragraphs = doc.paragraphs

    section = doc.sections[-1]
    expected = {
        "page_width": 11905,
        "page_height": 16838,
        "top_margin": 720,
        "bottom_margin": 677,
        "left_margin": 893,
        "right_margin": 893,
        "header_distance": 403,
        "footer_distance": 403,
    }
    for field, target in expected.items():
        actual = twips(getattr(section, field))
        if abs(actual - target) > 2:
            errors.append(f"正文节{field}={actual} twips，应为{target}±2")

    body_text = document_text(path, after_page_break=after_page_break)
    if not any(marker.lower() in body_text.lower() for marker in NAME_MARKERS):
        errors.append("简历正文未识别到姓名标题")
    section_headings = []
    for paragraph in body_paragraphs:
        if has_page_break(paragraph):
            errors.append(f"正文仍含显式分页：{paragraph.text[:30]}")
        style_name = paragraph.style.name if paragraph.style else ""
        if style_name in ("Heading 1", "SectionHeading"):
            section_headings.append(paragraph)
    if len(section_headings) < 3:
        errors.append(f"正文栏目标题过少：{len(section_headings)}")
    for paragraph in section_headings:
        p_pr = paragraph._p.pPr
        shd = p_pr.find(qn("w:shd")) if p_pr is not None else None
        fill = shd.get(qn("w:fill")) if shd is not None else None
        if fill != "EAF0F6":
            errors.append(f"栏目标题缺少EAF0F6底纹：{paragraph.text[:30]}")
    return errors


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("resume", type=Path)
    parser.add_argument("--baseline", type=Path)
    parser.add_argument("--after-page-break", action="store_true", help="只审计审阅区后的简历正文；支持显式分页或分节边界")
    parser.add_argument("--min-ratio", type=float, default=0.85)
    parser.add_argument("--require-section", action="append", default=[])
    parser.add_argument("--check-format", action="store_true")
    parser.add_argument("--require-font-installed", action="store_true")
    parser.add_argument("--context", type=Path, help="复用本岗位验收记录中的 resume_review，核对来源、项目归属和局部修改")
    parser.add_argument("--name-marker", action="append", default=[], help="Current candidate name; repeat for aliases")
    parser.add_argument("--case-marker", action="append", default=[], help="Current candidate project name; repeat when needed")
    args = parser.parse_args()
    global NAME_MARKERS, CASE_PATTERN
    NAME_MARKERS = tuple(args.name_marker)
    CASE_PATTERN = "|".join(re.escape(x) for x in args.case_marker) or r"(?!)"
    if (args.after_page_break or args.check_format) and not NAME_MARKERS:
        parser.error("--name-marker is required for divider or format checks")

    errors: list[str] = []
    try:
        text = document_text(args.resume, after_page_break=args.after_page_break)
    except (ValueError, KeyError, zipfile.BadZipFile) as exc:
        print("ERROR:", exc)
        return 1
    errors.extend(summary_errors(text))
    errors.extend(revision_errors(args.resume))
    if args.context:
        try:
            context = json.loads(args.context.read_text(encoding="utf-8-sig"))
            errors.extend(context_errors(text, context.get("resume_review", context)))
        except (OSError, ValueError, KeyError, TypeError) as exc:
            errors.append("无法核验本岗位修改记录:" + str(exc))
    size = effective_chars(text)
    if size < 300:
        errors.append(f"正文过短或提取失败：{size} 字符")
    hits = [label for label, pattern in FORBIDDEN_PATTERNS if re.search(pattern, text, flags=re.IGNORECASE)]
    if hits:
        errors.append("含内部判断/禁止词：" + "、".join(hits))
    missing_sections = [section for section in args.require_section if section not in text]
    if missing_sections:
        errors.append("缺少关键栏目：" + "、".join(missing_sections))

    ratio = None
    if args.baseline:
        try:
            baseline_text = document_text(args.baseline, after_page_break=args.after_page_break)
        except (ValueError, KeyError, zipfile.BadZipFile) as exc:
            print("ERROR: 无法读取基准稿：", exc, sep="")
            return 1
        baseline_size = effective_chars(baseline_text)
        if baseline_size == 0:
            errors.append("基准稿正文为空")
        else:
            ratio = size / baseline_size
            if ratio < args.min_ratio:
                errors.append(f"正文篇幅比 {ratio:.1%}，低于阈值 {args.min_ratio:.1%}")
    if args.require_font_installed and not font_installed():
        errors.append("本机未安装 Microsoft YaHei；停止交付")
    if args.check_format:
        errors.extend(format_errors(args.resume, after_page_break=args.after_page_break))

    print(f"resume_chars={size}")
    if ratio is not None:
        print(f"baseline_ratio={ratio:.4f}")
    if errors:
        for error in errors:
            print("ERROR:", error)
        return 1
    print("PASS: automated checks only; source-paragraph and rendered-page reviews remain required; user approval/application not inferred")
    return 0


if __name__ == "__main__":
    sys.exit(main())
